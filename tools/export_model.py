#!/usr/bin/env python3
"""Extract the stat model from the Total Resources calculator into data/model.json.

Starting from the 'Total' column of the Statmath sheet, follow every cell reference
transitively and dump the formulas (or constants) of every cell reached. The app's
tiny spreadsheet engine (js/engine.js) evaluates them. Google-Sheets-only wrappers
are unwrapped; SERIESSUM(x,0,1,INDEX(SIGN(ROW(INDIRECT("1:"&k))),,)) is rewritten as
GEOSUM(x,k) (= sum of x^i for i < k).

Usage: python3 tools/export_model.py <Obelisk_Total_Resources_Calculator.xlsx> data/model.json
"""
import json, re, sys
from collections import deque
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, column_index_from_string

WRAP = re.compile(r'^=IFERROR\(__xludf\.DUMMYFUNCTION\("(.*)"\),(.*)\)$', re.S)
REF = re.compile(r"(?:'([^']+)'|([A-Za-z_][A-Za-z0-9_ ]*?))!\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?|(?<![A-Za-z!'\d])\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?(?![A-Za-z(])")
SERIES = re.compile(r'SERIESSUM\(\s*([\d.]+)\s*,\s*0\s*,\s*1\s*,\s*INDEX\(SIGN\(ROW\(INDIRECT\("1:"&([^)]*?)\)\)\),,\)\)', re.I)

def unwrap(v):
    if hasattr(v, 'text'): v = '=' + str(v.text)
    s = str(v)
    m = WRAP.match(s)
    if m: s = '=' + m.group(1).replace('""', '"')
    return s

def norm(s):
    s = s.replace('\n', ' ')
    s = SERIES.sub(lambda m: f'GEOSUM({m.group(1)},{m.group(2)})', s)
    return s

def refs_in(formula, cur_sheet):
    out = []
    # strip string literals first
    f = re.sub(r'"[^"]*"', '""', formula)
    for m in REF.finditer(f):
        if m.group(3):
            sheet = m.group(1) or m.group(2); c1, r1, c2, r2 = m.group(3), m.group(4), m.group(5), m.group(6)
        else:
            sheet = cur_sheet; c1, r1, c2, r2 = m.group(7), m.group(8), m.group(9), m.group(10)
        sheet = sheet.strip()
        if c2:
            for ci in range(column_index_from_string(c1), column_index_from_string(c2) + 1):
                for ri in range(int(r1), int(r2) + 1):
                    out.append((sheet, f'{get_column_letter(ci)}{ri}'))
        else:
            out.append((sheet, f'{c1}{r1}'))
    return out

def main(src, dst):
    wb = load_workbook(src, read_only=True)
    sheets = {}
    def cell(sheet, addr):
        if sheet not in sheets:
            try: sheets[sheet] = {c.coordinate: c.value for row in wb[sheet].iter_rows() for c in row if c.value is not None}
            except KeyError: sheets[sheet] = {}
        return sheets[sheet].get(addr)
    sm = wb['Statmath']
    headers = [c.value for c in next(sm.iter_rows(min_row=1, max_row=1, min_col=1, max_col=22))]
    stats = []
    for row in sm.iter_rows(min_row=2, max_row=240, min_col=1, max_col=22):
        name, total = row[0].value, row[2].value
        if name and total is not None and isinstance(total, str) and total.startswith('='):
            parts = {}
            for i in range(3, 22):
                v = row[i].value
                if v is not None and headers[i]: parts[headers[i]] = f'Statmath!{row[i].coordinate}'
            stats.append({'name': name.strip(), 'cell': f'Statmath!{row[2].coordinate}', 'parts': parts})
    seen = {}; q = deque()
    for s in stats:
        q.append(tuple(s['cell'].split('!')))
        for pc in s['parts'].values(): q.append(tuple(pc.split('!')))
    missing = 0
    while q:
        sh, ad = q.popleft()
        key = f'{sh}!{ad}'
        if key in seen: continue
        v = cell(sh, ad)
        if v is None: seen[key] = {'v': None}; missing += 1; continue
        if isinstance(v, str) and (v.startswith('=') or hasattr(v, 'text')) or hasattr(v, 'text'):
            s = norm(unwrap(v))
            if 'COMPUTED_VALUE' in s: seen[key] = {'v': None}; continue
            seen[key] = {'f': s[1:]}
            for r in refs_in(s, sh): q.append(r)
        else:
            seen[key] = {'v': v}
    inputs = sorted(k for k, c in seen.items() if 'f' not in c)
    cells = {k: c for k, c in seen.items() if not ('v' in c and c['v'] is None)}  # blanks are implicit
    blanks = sorted(k for k, c in seen.items() if 'v' in c and c['v'] is None and not k.startswith('Statmath!'))
    with open(dst, 'w') as fh:
        json.dump({'_source': 'Obelisk Total Resources Calculator v7.2', 'stats': stats, 'cells': cells, 'blanks': blanks}, fh)
    print(f'{len(stats)} stats, {len(seen)} cells ({sum(1 for c in seen.values() if "f" in c)} formulas, {len(inputs)} inputs/constants, {missing} blank)')
    funcs = {}
    for c in seen.values():
        if 'f' in c:
            for fn in re.findall(r'([A-Za-z_]+)\(', c['f']): funcs[fn.upper()] = funcs.get(fn.upper(), 0) + 1
    print('functions:', sorted(funcs.items(), key=lambda x: -x[1]))

if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
