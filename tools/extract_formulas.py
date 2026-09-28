#!/usr/bin/env python3
"""Extract every formula from the community Google-Sheets calculators (xlsx exports)
into readable Markdown reference files, one per workbook.

Google Sheets wraps functions Excel cannot evaluate (ARRAYFORMULA, FILTER, LET,
REGEXEXTRACT, IMPORTRANGE ...) as  =IFERROR(__xludf.DUMMYFUNCTION("<original>"),<cached>)
The original formula text survives inside the string, so nothing is lost; this
script unwraps it. Cells that only hold a spilled cached value ("COMPUTED_VALUE")
are skipped because they carry no logic of their own.

Usage:  python3 tools/extract_formulas.py <folder-with-xlsx> docs/formulas
Needs:  pip install openpyxl
"""
import glob
import os
import re
import sys

from openpyxl import load_workbook

WRAP = re.compile(r'^=IFERROR\(__xludf\.DUMMYFUNCTION\("(.*)"\),(.*)\)$', re.S)


def unwrap(formula):
    m = WRAP.match(formula)
    if not m:
        return formula, False
    inner = m.group(1).replace('""', '"')
    return '=' + inner, True


def main(src, dst):
    os.makedirs(dst, exist_ok=True)
    index = ['# Formula reference', '',
             'Extracted from the community calculators. One file per workbook.',
             'Rows tagged *(Sheets-only)* used Google-Sheets functions that Excel',
             'cannot evaluate; the original formula text is shown.', '']
    for path in sorted(glob.glob(os.path.join(src, '*.xlsx'))):
        name = os.path.splitext(os.path.basename(path))[0]
        wb = load_workbook(path, read_only=True)
        out = [f'# {name}', '']
        total = 0
        for ws in wb.worksheets:
            rows = []
            for row in ws.iter_rows():
                for c in row:
                    v = c.value
                    if isinstance(v, str) and v.startswith('='):
                        if 'COMPUTED_VALUE' in v:
                            continue
                        f, was_sheets = unwrap(v)
                        rows.append((c.coordinate, f, was_sheets))
                    elif hasattr(v, 'text'):  # openpyxl ArrayFormula
                        rows.append((c.coordinate, '=' + str(v.text), False))
            if not rows:
                continue
            total += len(rows)
            out += [f'## Sheet: {ws.title}', '', f'{len(rows)} formulas', '',
                    '| Cell | Formula |', '|---|---|']
            for coord, f, was_sheets in rows:
                f = f.replace('|', '\\|').replace('\n', ' ')
                tag = ' *(Sheets-only)*' if was_sheets else ''
                out.append(f'| {coord} | `{f}`{tag} |')
            out.append('')
        with open(os.path.join(dst, name + '.md'), 'w') as fh:
            fh.write('\n'.join(out))
        index.append(f'- [{name}]({name}.md) ({total} formulas)')
        print(f'{name}: {total} formulas')
    with open(os.path.join(dst, 'README.md'), 'w') as fh:
        fh.write('\n'.join(index) + '\n')


if __name__ == '__main__':
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)
    main(sys.argv[1], sys.argv[2])
